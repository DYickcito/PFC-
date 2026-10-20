"""
Servicio RAG principal - Orquestado con LlamaIndex 0.14.x.
Gestiona:
  - Extracción de texto por tipo de archivo (PDF, DOCX, XLSX, CSV, TXT)
  - Limpieza del texto extraído antes de segmentarlo
  - Embeddings con Google Gemini (models/gemini-embedding-2, 3072 dims)
  - Rate-limiting en ingesta para no superar cuota gratuita de Gemini
  - LLM dinámico con Groq (selección de modelo por solicitud)
  - Enrutamiento por puntaje de similitud: si ningún fragmento supera el
    umbral, se responde con conocimiento general y se avisa al usuario
  - Atajo para saludos y preguntas sobre el asistente (no busca en Qdrant)
  - Caché en memoria de respuestas recientes y respuesta en streaming
"""
import asyncio
import csv
import json
import logging
import math
import re
import time
import unicodedata
from collections import Counter, OrderedDict
from functools import lru_cache
from pathlib import Path
from typing import AsyncIterator, List, Optional

from llama_index.core import (
    PromptTemplate,
    Settings as LlamaSettings,
    StorageContext,
    VectorStoreIndex,
)
from llama_index.core.node_parser import SentenceSplitter
from llama_index.core.schema import Document, MetadataMode, NodeWithScore, QueryBundle
from llama_index.embeddings.gemini import GeminiEmbedding
from llama_index.llms.groq import Groq
from llama_index.vector_stores.qdrant import QdrantVectorStore

from app.core.config import settings
from app.db.qdrant_setup import qdrant_client

logger = logging.getLogger(__name__)

# ─────────────────────────────────────────────
# Constantes de rate-limiting (Gemini Free Tier)
# ~5 req/min → esperamos 12s entre lotes de 5 nodos
# ─────────────────────────────────────────────
BATCH_SIZE = 5
BATCH_DELAY_SEC = 12.0

# ─────────────────────────────────────────────
# Catálogo único de modelos LLM disponibles en Groq.
# Es la única fuente de verdad: el router y el frontend leen de aquí.
# extra_kwargs se pasa al cliente de Groq; los parámetros propios de Groq
# van dentro de extra_body porque el cliente openai no los acepta sueltos.
# Qwen 3.6 27B se retiró del catálogo porque no respondía en las pruebas.
# ─────────────────────────────────────────────
#
# context_window: tamaño real de contexto del modelo. Si no se indica,
# LlamaIndex asume 3 900 tokens y, al restar la reserva de salida, parte los
# fragmentos en varias llamadas al LLM (más tokens y más errores 429).
# reasoning_effort "low": GPT-OSS razona menos antes de responder; para
# preguntas sobre documentos basta y se gastan menos tokens de salida.
LLM_MODELS = {
    "openai/gpt-oss-20b": {
        "label": "GPT-OSS 20B (Groq)",
        "context_window": 131072,
        "extra_kwargs": {"extra_body": {"reasoning_effort": "low"}},
    },
}
MAX_OUTPUT_TOKENS = 1024
DEFAULT_MODEL = "openai/gpt-oss-20b"
ALLOWED_MODELS = set(LLM_MODELS)

# ─────────────────────────────────────────────
# Prompts
# ─────────────────────────────────────────────
QA_PROMPT = PromptTemplate(
    "Eres el asistente académico de la carrera de Diseño y Desarrollo de "
    "Software de Tecsup. Responde en español usando únicamente la información "
    "de los fragmentos de documentos oficiales que aparecen abajo.\n"
    "Si los fragmentos no alcanzan para responder la pregunta completa, dilo "
    "con claridad en lugar de completar con suposiciones.\n"
    "---------------------\n"
    "{context_str}\n"
    "---------------------\n"
    "Pregunta: {query_str}\n"
    "Respuesta:"
)

FALLBACK_PROMPT = (
    "Eres el asistente académico de la carrera de Diseño y Desarrollo de "
    "Software de Tecsup. Respondes consultas a partir de estos documentos "
    "cargados en el sistema:\n"
    "{catalogo}\n\n"
    "Para la pregunta del estudiante no se encontraron fragmentos relevantes en "
    "esos documentos. Responde en español siguiendo estas reglas:\n"
    "1. Si el estudiante saluda, pregunta qué datos tienes, qué sabes o qué puedes "
    "hacer, preséntate en una o dos frases, explica qué temas puedes responder "
    "según la lista de documentos y da dos o tres ejemplos de preguntas que puede "
    "hacerte. En este caso no digas que la pregunta está fuera de tu información.\n"
    "2. En cualquier otro caso, empieza con una sola frase breve y amable, con tus "
    "propias palabras, que diga que la pregunta no está relacionada con la "
    "información académica que manejas. Por ejemplo: \"Tu pregunta no tiene "
    "relación con la información académica que manejo, pero te cuento lo que "
    "sé:\". No repitas siempre la misma frase.\n"
    "3. Si es una pregunta general (cultura general, programación, conceptos), "
    "respóndela de forma breve con tu conocimiento general.\n"
    "4. Si pide datos propios de Tecsup o de otra institución (horarios, docentes, "
    "costos, trámites, servicios, fechas) que no aparecen en la lista de "
    "documentos, no inventes nada: di que no tienes esa información y recomienda "
    "consultarla con la coordinación de la carrera.\n\n"
    "Pregunta: {query}\n"
    "Respuesta:"
)


# ─────────────────────────────────────────────
# Inicialización del embedding model en LlamaIndex Settings global
# ─────────────────────────────────────────────
def init_rag_settings() -> None:
    """
    Configura el modelo de embeddings Gemini como predeterminado global
    de LlamaIndex. Se llama UNA SOLA VEZ desde el lifespan de FastAPI (main.py).
    El LLM NO se fija de forma global: se crea en cada consulta y se pasa
    explícitamente, para que dos usuarios con modelos distintos no se pisen.
    """
    LlamaSettings.embed_model = GeminiEmbedding(
        model_name="models/gemini-embedding-2",
        api_key=settings.gemini_api_key,
    )
    LlamaSettings.llm = None
    logger.info("✅ Gemini embedding model inicializado: models/gemini-embedding-2")


# ─────────────────────────────────────────────
# Limpieza de texto
# ─────────────────────────────────────────────
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f​-‏﻿]")
_PAGE_NUMBER = re.compile(r"^(p[áa]g(ina)?\.?\s*)?\d{1,4}(\s*(de|/)\s*\d{1,4})?$", re.IGNORECASE)
_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")
_MULTI_SPACE = re.compile(r"[ \t ]{2,}")
_MULTI_NEWLINE = re.compile(r"\n{3,}")


_BORDE = 2  # líneas al inicio y al final de cada página donde se buscan encabezados y pies


def _normalizar(texto: str) -> str:
    texto = unicodedata.normalize("NFKC", texto)
    texto = _CONTROL_CHARS.sub("", texto)
    return texto.replace("\r\n", "\n").replace("\r", "\n")


def _indices_borde(lineas: List[str]) -> List[int]:
    """Índices de las primeras y últimas líneas con texto de una página."""
    con_texto = [i for i, ln in enumerate(lineas) if ln.strip()]
    return sorted(set(con_texto[:_BORDE] + con_texto[-_BORDE:]))


def _repeated_page_lines(pages: List[str]) -> set:
    """
    Detecta encabezados y pies de página: líneas cortas que aparecen en el borde
    (primeras o últimas líneas) de al menos el 60 % de las páginas de un PDF.
    Solo se miran los bordes para no borrar contenido que se repite dentro del
    texto, como celdas de tablas o nombres de cursos.
    """
    if len(pages) < 2:
        return set()
    counts = Counter()
    for page in pages:
        lineas = page.split("\n")
        bordes = {lineas[i].strip() for i in _indices_borde(lineas) if len(lineas[i].strip()) <= 100}
        counts.update(bordes)
    limite = max(2, math.ceil(0.6 * len(pages)))
    return {linea for linea, n in counts.items() if n >= limite}


def _limpiar_bordes(page: str, repetidas: set) -> str:
    """Quita de los bordes de una página los encabezados repetidos y los números de página."""
    lineas = page.split("\n")
    quitar = {
        i for i in _indices_borde(lineas)
        if lineas[i].strip() in repetidas or _PAGE_NUMBER.match(lineas[i].strip())
    }
    return "\n".join(ln for i, ln in enumerate(lineas) if i not in quitar)


def clean_text(text: str, pages: Optional[List[str]] = None) -> str:
    """
    Limpia el texto extraído antes de segmentarlo:
      1. Normaliza unicode (NFKC) y elimina caracteres de control o invisibles.
      2. Solo en PDF: quita encabezados y pies repetidos y números de página,
         buscándolos únicamente en las primeras y últimas líneas de cada página
         (así no se borran números ni textos repetidos dentro de las tablas).
      3. Une palabras cortadas con guion al final de línea.
      4. Colapsa espacios repetidos y saltos de línea excesivos.
    Si se pasan las páginas, el texto se reconstruye a partir de ellas.
    """
    if pages:
        paginas = [_normalizar(p) for p in pages]
        repetidas = _repeated_page_lines(paginas)
        text = "\n\n".join(_limpiar_bordes(p, repetidas) for p in paginas)
    else:
        text = _normalizar(text)

    text = "\n".join(ln.rstrip() for ln in text.split("\n"))
    text = _HYPHEN_BREAK.sub(r"\1\2", text)
    text = _MULTI_SPACE.sub(" ", text)
    text = _MULTI_NEWLINE.sub("\n\n", text)
    return text.strip()


# ─────────────────────────────────────────────
# Extracción de texto por tipo de archivo
# ─────────────────────────────────────────────
def _extract_text_from_file(file_path: str) -> str:
    """
    Extrae el texto plano de un archivo según su extensión y lo limpia.
    Parsers utilizados:
      .pdf   → pypdf
      .docx  → python-docx
      .xlsx  → openpyxl
      .csv   → csv (stdlib)
      .txt   → open() directo
    """
    path = Path(file_path)
    ext = path.suffix.lower()
    logger.info(f"📂 Extrayendo texto de '{path.name}' (tipo: {ext})")

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        pages = [page.extract_text() or "" for page in reader.pages]
        pages = [p for p in pages if p.strip()]
        logger.info(f"   PDF: {len(reader.pages)} páginas extraídas.")
        return clean_text("\n\n".join(pages), pages=pages)

    elif ext == ".docx":
        from docx import Document as DocxDocument
        doc = DocxDocument(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # También extraer texto de tablas
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        logger.info(f"   DOCX: {len(paragraphs)} párrafos/filas extraídos.")
        return clean_text("\n".join(paragraphs))

    elif ext in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        rows_text = []
        for sheet in wb.worksheets:
            rows_text.append(f"[Hoja: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                row_str = " | ".join(str(v) for v in row if v is not None)
                if row_str.strip():
                    rows_text.append(row_str)
        logger.info(f"   XLSX: {len(rows_text)} filas extraídas.")
        return clean_text("\n".join(rows_text))

    elif ext == ".csv":
        rows_text = []
        with open(file_path, newline="", encoding="utf-8-sig", errors="ignore") as f:
            reader = csv.reader(f)
            for row in reader:
                row_str = " | ".join(cell for cell in row if cell.strip())
                if row_str:
                    rows_text.append(row_str)
        logger.info(f"   CSV: {len(rows_text)} filas extraídas.")
        return clean_text("\n".join(rows_text))

    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        logger.info(f"   TXT: {len(text)} caracteres extraídos.")
        return clean_text(text)

    else:
        raise ValueError(f"Tipo de archivo no soportado para extracción: {ext}")


# ─────────────────────────────────────────────
# Helpers de Vector Store / Index
# ─────────────────────────────────────────────
def _get_vector_store() -> QdrantVectorStore:
    return QdrantVectorStore(
        client=qdrant_client,
        collection_name=settings.qdrant_collection_name,
    )


def _get_storage_context() -> StorageContext:
    return StorageContext.from_defaults(vector_store=_get_vector_store())


@lru_cache(maxsize=1)
def _load_index_from_qdrant() -> VectorStoreIndex:
    """
    Carga el índice existente desde Qdrant sin re-indexar. Se crea una sola
    vez y se reutiliza: el índice no guarda datos propios, cada búsqueda va a
    Qdrant, así que los documentos nuevos o eliminados se ven al instante.
    """
    return VectorStoreIndex.from_vector_store(
        vector_store=_get_vector_store(),
    )


# ─────────────────────────────────────────────
# Ingesta de Documentos con Rate-Limiting
# ─────────────────────────────────────────────
async def ingest_document(
    file_path: str,
    document_id: str,
    document_metadata: Optional[dict] = None,
) -> int:
    """
    Carga, limpia, segmenta e indexa un documento en Qdrant.
    Pipeline:
      1. Extracción y limpieza de texto (por tipo de archivo)
      2. Segmentación en chunks (SentenceSplitter)
      3. Vectorización + indexación en lotes con pausa (rate-limit Gemini)

    Las llamadas a Gemini y Qdrant son bloqueantes, así que se ejecutan en
    un hilo aparte para no congelar el servidor mientras se indexa.
    Retorna el número de chunks (nodos) indexados.
    """
    logger.info(f"📄 Iniciando ingesta: {file_path}")

    # 1. Extraer y limpiar texto
    raw_text = await asyncio.to_thread(_extract_text_from_file, file_path)

    if not raw_text.strip():
        raise ValueError(f"No se pudo extraer texto del archivo: {file_path}")

    # 2. Crear documento LlamaIndex con metadatos enriquecidos
    # id_ = ID de Supabase: LlamaIndex lo guarda en Qdrant como document_id /
    # ref_doc_id, así las fuentes citadas apuntan al registro real y el
    # documento se puede borrar de Qdrant por ese mismo ID.
    base_metadata = dict(document_metadata or {})
    base_metadata["document_id"] = document_id
    document = Document(id_=document_id, text=raw_text, metadata=base_metadata)

    # 3. Parsear en chunks
    parser = SentenceSplitter(chunk_size=512, chunk_overlap=64)
    nodes = parser.get_nodes_from_documents([document])
    total_nodes = len(nodes)
    logger.info(f"📦 {total_nodes} chunks generados de '{Path(file_path).name}'")

    # 4. Indexar en lotes con pausa (rate-limiting Gemini)
    storage_context = _get_storage_context()
    indexed_count = 0
    total_batches = (total_nodes + BATCH_SIZE - 1) // BATCH_SIZE

    for batch_num, batch_start in enumerate(range(0, total_nodes, BATCH_SIZE), start=1):
        batch = nodes[batch_start: batch_start + BATCH_SIZE]
        logger.info(f"🔄 Lote {batch_num}/{total_batches} — {len(batch)} nodos...")

        await asyncio.to_thread(
            VectorStoreIndex,
            nodes=batch,
            storage_context=storage_context,
            show_progress=False,
        )
        indexed_count += len(batch)

        if batch_start + BATCH_SIZE < total_nodes:
            logger.info(f"⏳ Esperando {BATCH_DELAY_SEC}s (rate-limit Gemini)...")
            await asyncio.sleep(BATCH_DELAY_SEC)

    logger.info(f"✅ Ingesta finalizada: {indexed_count} nodos en Qdrant.")
    return indexed_count


def update_document_metadata(document_id: str, cambios: dict) -> int:
    """
    Actualiza metadatos (por ejemplo nombre_documento o id_carrera) en todos los
    fragmentos de un documento en Qdrant. LlamaIndex guarda los metadatos dos
    veces: como campos sueltos del payload y dentro de _node_content, que es de
    donde reconstruye las fuentes citadas; por eso se actualizan ambos.
    Retorna la cantidad de fragmentos actualizados.
    """
    from qdrant_client.models import FieldCondition, Filter, MatchValue

    filtro = Filter(must=[FieldCondition(key="document_id", match=MatchValue(value=document_id))])
    actualizados = 0
    offset = None
    while True:
        puntos, offset = qdrant_client.scroll(
            collection_name=settings.qdrant_collection_name,
            scroll_filter=filtro,
            with_payload=True,
            with_vectors=False,
            limit=256,
            offset=offset,
        )
        for punto in puntos:
            payload = dict(cambios)
            contenido = punto.payload.get("_node_content")
            if contenido:
                nodo = json.loads(contenido)
                nodo.setdefault("metadata", {}).update(cambios)
                payload["_node_content"] = json.dumps(nodo)
            qdrant_client.set_payload(
                collection_name=settings.qdrant_collection_name,
                payload=payload,
                points=[punto.id],
            )
            actualizados += 1
        if offset is None:
            break
    return actualizados


def delete_document_vectors(document_id: str) -> None:
    """Elimina de Qdrant los vectores de un documento (útil si la ingesta falla a medias)."""
    _get_vector_store().delete(ref_doc_id=document_id)


# ─────────────────────────────────────────────
# LLM Groq — selección dinámica por request
# ─────────────────────────────────────────────
@lru_cache(maxsize=None)
def _build_groq_llm(model_name: str) -> Groq:
    """
    Instancia el LLM Groq con el modelo especificado y validado. Se guarda
    uno por modelo para reutilizar la conexión HTTPS entre consultas.
    """
    if model_name not in ALLOWED_MODELS:
        raise ValueError(
            f"Modelo '{model_name}' no permitido. "
            f"Opciones válidas: {sorted(ALLOWED_MODELS)}"
        )
    info = LLM_MODELS[model_name]
    return Groq(
        model=model_name,
        api_key=settings.groq_api_key,
        temperature=0.2,
        max_tokens=MAX_OUTPUT_TOKENS,
        context_window=info.get("context_window", 8192),
        additional_kwargs=info["extra_kwargs"],
    )


def _format_sources(nodes: List[NodeWithScore]) -> list:
    return [
        {
            "document_id": node.metadata.get("document_id", "N/A"),
            "nombre_documento": node.metadata.get("nombre_documento", "N/A"),
            "score": round(node.score, 4) if node.score is not None else None,
            "texto_fragmento": node.get_content()[:300],
        }
        for node in nodes
    ]


def route_by_similarity(nodes: List[NodeWithScore], threshold: float):
    """
    Enrutamiento por puntaje de similitud.
    Retorna (nodos_relevantes, score_maximo). Si ningún nodo alcanza el
    umbral, la lista viene vacía y la consulta se responde sin contexto.
    """
    scores = [n.score for n in nodes if n.score is not None]
    max_score = max(scores) if scores else None
    relevant = [n for n in nodes if n.score is not None and n.score >= threshold]
    return relevant, max_score


# ─────────────────────────────────────────────
# Consulta RAG
# ─────────────────────────────────────────────
# ─────────────────────────────────────────────
# Catálogo de documentos cargados (para saludos y preguntas sobre el asistente)
# ─────────────────────────────────────────────
# Se invalida al subir, editar o eliminar documentos, así que puede durar una hora
_CATALOGO_TTL_SEC = 3600
_catalogo_cache = {"momento": 0.0, "texto": ""}


def invalidar_catalogo() -> None:
    """
    Fuerza a releer el catálogo en la próxima consulta (tras subir, editar o
    eliminar). También vacía la caché de respuestas, porque una respuesta
    guardada podría citar un documento que cambió.
    """
    _catalogo_cache["momento"] = 0.0
    _respuestas_cache.clear()


def catalogo_documentos() -> str:
    """
    Lista de documentos registrados en Supabase (nombre, tipo, carrera y
    descripción), en texto para el prompt. Se guarda en memoria (una hora o
    hasta que cambien los documentos) para no consultar Supabase en cada pregunta. Si la consulta falla, se usa
    el último valor conocido.
    """
    ahora = time.monotonic()
    if _catalogo_cache["texto"] and ahora - _catalogo_cache["momento"] < _CATALOGO_TTL_SEC:
        return _catalogo_cache["texto"]
    try:
        from app.db.supabase_client import get_supabase_client

        filas = (
            get_supabase_client()
            .table("documento")
            .select("nombre_documento, descripcion, tipo_documento ( nombre_tipo ), carrera ( nombre_carrera )")
            .order("fecha_carga", desc=True)
            .limit(50)
            .execute()
        ).data or []
        lineas = []
        for f in filas:
            tipo = (f.get("tipo_documento") or {}).get("nombre_tipo", "")
            carrera = (f.get("carrera") or {}).get("nombre_carrera", "")
            linea = f"- {f.get('nombre_documento', 'Documento')} ({tipo}, {carrera})"
            if f.get("descripcion"):
                linea += f": {f['descripcion']}"
            lineas.append(linea)
        texto = "\n".join(lineas) if lineas else "- (todavía no hay documentos cargados)"
        _catalogo_cache.update(momento=ahora, texto=texto)
        return texto
    except Exception as e:
        logger.warning(f"⚠️ No se pudo leer el catálogo de documentos: {e}")
        return _catalogo_cache["texto"] or "- (no se pudo obtener la lista de documentos)"


# ─────────────────────────────────────────────
# Atajo para saludos y preguntas sobre el asistente
# ─────────────────────────────────────────────
_PALABRAS_SALUDO = {
    "hola", "holi", "holaa", "buenas", "buenos", "buen", "dia", "dias", "tardes", "noches",
    "hey", "saludos", "que", "tal", "como", "estas", "gracias", "muchas", "mil", "ok", "okay",
    "vale", "chau", "chao", "adios", "hasta", "luego", "pronto", "nos", "vemos", "bot", "asistente",
}
_PREGUNTA_META = re.compile(
    r"^(que|cuales) (datos|informacion|documentos|temas|archivos) (tienes|manejas|conoces|hay|sabes)"
    r"|^que (sabes|puedes) (hacer|responder|contestar)"
    r"|^(quien|que) eres|^como te llamas|^para que sirves|^(ayuda|ayudame|help)$"
)


def _normalizar_pregunta(texto: str) -> str:
    """Minúsculas, sin tildes ni signos de puntuación y con espacios simples."""
    texto = unicodedata.normalize("NFKD", texto.lower())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"[^\w\s]", " ", texto)
    return " ".join(texto.split())


def es_conversacional(prompt: str) -> bool:
    """
    True si el mensaje es solo un saludo, un agradecimiento o una pregunta
    sobre el propio asistente ("hola", "gracias", "qué datos tienes"). Esos
    mensajes no necesitan buscar en Qdrant: se responden directo con el
    catálogo de documentos, lo que ahorra el embedding y la búsqueda.
    """
    texto = _normalizar_pregunta(prompt)
    palabras = texto.split()
    if not palabras or len(palabras) > 8:
        return False
    return all(p in _PALABRAS_SALUDO for p in palabras) or bool(_PREGUNTA_META.search(texto))


# ─────────────────────────────────────────────
# Caché de respuestas recientes
# Muchos alumnos hacen la misma pregunta; si ya se respondió hace poco y los
# documentos no cambiaron, se devuelve la misma respuesta sin llamar a Gemini,
# Qdrant ni Groq. Vive en memoria: se pierde al reiniciar el backend.
# ─────────────────────────────────────────────
_CACHE_TTL_SEC = 600
_CACHE_MAX = 200
_respuestas_cache: "OrderedDict[tuple, tuple[float, dict]]" = OrderedDict()


def _cache_key(prompt: str, model_name: str, similarity_top_k: int) -> tuple:
    return (_normalizar_pregunta(prompt), model_name, similarity_top_k)


def _cache_get(key: tuple) -> Optional[dict]:
    item = _respuestas_cache.get(key)
    if not item:
        return None
    momento, resultado = item
    if time.monotonic() - momento > _CACHE_TTL_SEC:
        _respuestas_cache.pop(key, None)
        return None
    _respuestas_cache.move_to_end(key)
    return resultado


def _cache_put(key: tuple, resultado: dict) -> None:
    _respuestas_cache[key] = (time.monotonic(), resultado)
    _respuestas_cache.move_to_end(key)
    while len(_respuestas_cache) > _CACHE_MAX:
        _respuestas_cache.popitem(last=False)


# ─────────────────────────────────────────────
# Consulta: preparación compartida por la versión normal y la de streaming
# ─────────────────────────────────────────────
def _ms(desde: float) -> int:
    return int((time.perf_counter() - desde) * 1000)


def _preparar_consulta(prompt: str, similarity_top_k: int) -> dict:
    """
    Parte bloqueante previa al LLM (se ejecuta en un hilo aparte):
    recupera fragmentos, enruta por puntaje y arma el prompt final.
    """
    threshold = settings.similarity_threshold
    t0 = time.perf_counter()

    if es_conversacional(prompt):
        logger.info("💬 Mensaje conversacional: se responde sin buscar en Qdrant")
        return {
            "llm_prompt": FALLBACK_PROMPT.format(catalogo=catalogo_documentos(), query=prompt),
            "sources": [],
            "uso_contexto": False,
            "conversacional": True,
            "max_score": None,
            "umbral": threshold,
            "recuperacion_ms": _ms(t0),
            "embedding_ms": 0,
            "busqueda_ms": 0,
        }

    # 1. Recuperar los fragmentos más parecidos: primero el vector de la
    #    pregunta (Gemini) y luego la búsqueda (Qdrant), medidos por separado
    embedding = LlamaSettings.embed_model.get_query_embedding(prompt)
    embedding_ms = _ms(t0)
    t1 = time.perf_counter()
    retriever = _load_index_from_qdrant().as_retriever(similarity_top_k=similarity_top_k)
    retrieved = retriever.retrieve(QueryBundle(query_str=prompt, embedding=embedding))
    busqueda_ms = _ms(t1)

    # 2. Enrutar según el puntaje
    relevant, max_score = route_by_similarity(retrieved, threshold)
    logger.info(
        f"📊 Score máximo={max_score} | umbral={threshold} | "
        f"fragmentos sobre el umbral={len(relevant)}/{len(retrieved)}"
    )

    if relevant:
        # 3a. Hay contexto oficial: el prompt lleva los fragmentos
        contexto = "\n\n".join(n.node.get_content(metadata_mode=MetadataMode.LLM) for n in relevant)
        llm_prompt = QA_PROMPT.format(context_str=contexto, query_str=prompt)
        sources = _format_sources(relevant)
    else:
        # 3b. Sin contexto suficiente: el modelo avisa con sus palabras que la
        #     pregunta está fuera de los documentos y responde con conocimiento
        #     general. La interfaz además lo marca con uso_contexto=False.
        llm_prompt = FALLBACK_PROMPT.format(catalogo=catalogo_documentos(), query=prompt)
        sources = []

    return {
        "llm_prompt": llm_prompt,
        "sources": sources,
        "uso_contexto": bool(relevant),
        "conversacional": False,
        "max_score": max_score,
        "umbral": threshold,
        "recuperacion_ms": _ms(t0),
        "embedding_ms": embedding_ms,
        "busqueda_ms": busqueda_ms,
    }


def _resultado(prep: dict, texto: str, model_name: str, generacion_ms: int) -> dict:
    return {
        "response": texto.strip(),
        "sources": prep["sources"],
        "model_used": model_name,
        "uso_contexto": prep["uso_contexto"],
        "conversacional": prep["conversacional"],
        "max_score": prep["max_score"],
        "umbral": prep["umbral"],
        "cache": False,
        "tiempos": {
            "recuperacion_ms": prep["recuperacion_ms"],
            "embedding_ms": prep["embedding_ms"],
            "busqueda_ms": prep["busqueda_ms"],
            "generacion_ms": generacion_ms,
        },
    }


def _desde_cache(resultado: dict) -> dict:
    return {**resultado, "cache": True, "tiempos": {"recuperacion_ms": 0, "generacion_ms": 0}}


def _run_query(prompt: str, model_name: str, similarity_top_k: int) -> dict:
    """Consulta completa sin streaming; se ejecuta en un hilo aparte."""
    llm = _build_groq_llm(model_name)
    key = _cache_key(prompt, model_name, similarity_top_k)
    guardado = _cache_get(key)
    if guardado:
        logger.info("⚡ Respuesta tomada de la caché")
        return _desde_cache(guardado)

    prep = _preparar_consulta(prompt, similarity_top_k)
    t0 = time.perf_counter()
    texto = llm.complete(prep["llm_prompt"]).text
    resultado = _resultado(prep, texto, model_name, _ms(t0))
    _cache_put(key, resultado)
    return resultado


async def query_rag(
    prompt: str,
    model_name: str,
    similarity_top_k: int = 5,
) -> dict:
    """
    Pipeline RAG completo:
      1. Si es un saludo o una pregunta sobre el asistente, no busca en Qdrant.
      2. Si no, recupera chunks de Qdrant usando embeddings Gemini y compara
         el score máximo con el umbral configurado.
      3. Si supera el umbral, genera la respuesta con esos fragmentos;
         si no, responde con conocimiento general y lo indica al usuario.

    Retorna: response, sources, model_used, uso_contexto, conversacional,
    max_score, umbral, cache y tiempos (recuperación y generación en ms).
    """
    logger.info(f"🔍 RAG Query | model={model_name} | prompt='{prompt[:80]}...'")
    return await asyncio.to_thread(_run_query, prompt, model_name, similarity_top_k)


async def query_rag_stream(
    prompt: str,
    model_name: str,
    similarity_top_k: int = 5,
) -> AsyncIterator[dict]:
    """
    Igual que query_rag, pero entrega la respuesta por partes para que la
    interfaz la vaya mostrando mientras el modelo escribe. Emite eventos:
      {"type": "meta", ...}   origen, fuentes y puntajes (antes del texto)
      {"type": "delta", "text"} cada trozo de texto generado
      {"type": "done", "resultado"} resultado completo, igual al de query_rag
    """
    logger.info(f"🔍 RAG Query (stream) | model={model_name} | prompt='{prompt[:80]}...'")
    llm = _build_groq_llm(model_name)
    key = _cache_key(prompt, model_name, similarity_top_k)
    guardado = _cache_get(key)

    if guardado:
        logger.info("⚡ Respuesta tomada de la caché")
        resultado = _desde_cache(guardado)
        yield {"type": "meta", **{k: v for k, v in resultado.items() if k != "response"}}
        yield {"type": "delta", "text": resultado["response"]}
        yield {"type": "done", "resultado": resultado}
        return

    prep = await asyncio.to_thread(_preparar_consulta, prompt, similarity_top_k)
    yield {
        "type": "meta",
        "sources": prep["sources"],
        "model_used": model_name,
        "uso_contexto": prep["uso_contexto"],
        "conversacional": prep["conversacional"],
        "max_score": prep["max_score"],
        "umbral": prep["umbral"],
        "cache": False,
    }

    t0 = time.perf_counter()
    partes: List[str] = []
    async for chunk in await llm.astream_complete(prep["llm_prompt"]):
        if chunk.delta:
            partes.append(chunk.delta)
            yield {"type": "delta", "text": chunk.delta}

    resultado = _resultado(prep, "".join(partes), model_name, _ms(t0))
    _cache_put(key, resultado)
    yield {"type": "done", "resultado": resultado}
